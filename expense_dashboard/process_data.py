# -*- coding: utf-8 -*-
"""
process_data.py — 第一階段：把上萬份「辨識等級(高)」轉檔結果 Excel 彙整進 DuckDB

【做什麼】
  1. 掃描指定資料夾 (可多個) 底下所有「XXX-辨識等級(高)」資料夾裡的 Excel
     (轉檔工具的輸出)，略過轉檔總表與 Excel 暫存檔 (~$)。
  2. 多個 process 平行讀檔 (python-calamine，沒有的話退回 openpyxl 唯讀
     模式)，每份檔案讀完只回傳清洗後的明細列；主程式每累積 BATCH_ROWS 列
     就寫進資料庫一次，記憶體用量固定，不會因為檔案多而爆掉。
  3. 清洗 (expense_normalize.py)：發票日期 -> 標準日期 + 年/月；AMOUNT/VAT
     -> 數值；Description -> 標準費用項目 + 費用大類；港口統一大小寫。
  4. 寫進 DuckDB (預設 data/expenses.duckdb)：
       expense_items  清洗後的完整明細 (原始費用名稱也保留在 description_raw)
       etl_files      每份 Excel 的處理紀錄 (修改時間、大小、列數、錯誤原因)
  5. 增量更新：再跑一次只處理新增或有修改的檔案；已經不存在的檔案，它的
     明細會從資料庫移除。加 --full 則整個重建。
  6. 同一張發票 (供應商 + 發票號碼) 出現在好幾份 Excel 時 (例如同一份 PDF
     轉了兩次)，只保留其中一份，其餘標記 is_duplicate，儀表板預設不計入。
  7. 讀不了/格式不符的 Excel 記進 etl_files 與 logs/expense_etl.log，
     不中斷整批處理。
  8. 每列記錄金額幣別 (currency)，最後用資料庫旁邊最新的作帳匯率檔更新匯率
     (fx.py)，儀表板用它換算成 USD。讀不了匯率檔只提示，不影響資料匯入。

儀表板「資料管理」頁 (data_admin.py) 也是呼叫這裡的 run()。

【怎麼用】
  python process_data.py "D:\\發票\\02-Maersk Logistics Indonesia"
  python process_data.py <資料夾1> <資料夾2> --supplier-folder MAERSK
  python process_data.py <資料夾> --full --parquet data\\expenses.parquet
"""

import argparse
import logging
import multiprocessing
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from datetime import datetime
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Callable, Dict, Iterator, List, Optional, Tuple

import duckdb
import pandas as pd

import fx
from expense_normalize import (
    DEFAULT_CURRENCY,
    SUPPLIER_NO_OWN_INVOICE,
    clean_text,
    detect_currency,
    folder_currency,
    is_no_own_invoice,
    normalize_expense,
    normalize_port,
    parse_amount,
    parse_date,
)
from settings import LOG_DIR, get_db_path

LOG_PATH = LOG_DIR / "expense_etl.log"

# 每累積幾列寫進資料庫一次 (控制主程式記憶體用量)
BATCH_ROWS = 20_000
# 標題列最多往下找幾列 (pdf_to_excel 範本在第 1 列、批次範本在第 2 列)
HEADER_SEARCH_ROWS = 6
SUMMARY_REPORT_PREFIX = "轉檔總表"
RESULT_SUFFIX = "_轉檔結果.xlsx"

# Excel 標題第一行 (換行前、轉大寫、去掉結尾的點) -> 欄位代碼
HEADER_FIELDS = {
    "INVOICE DATE": "invoice_date", "INVOICE NO": "invoice_no",
    "SUPPLIER": "supplier", "CONSIGNEE": "consignee", "DESCRIPTION": "description",
    "AMOUNT": "amount", "VAT": "vat", "ORDER NO": "order_no",
    "ARRIVE DATE": "arrive_date", "ON BOARD DATE": "onboard_date",
    "B/L NO": "bl_no", "MBL NO": "mbl_no", "PORT OF LOADING": "port_of_loading",
    "PORT OF DISCHARGE": "port_of_discharge", "VOLUME": "volume",
    "VESSEL": "vessel", "CONTAINER NO": "container_no", "SHIPPER": "shipper",
}

# expense_items 的欄位 (順序 = 建表順序 = 寫入 DataFrame 的欄位順序)
ITEM_COLUMNS = (
    ("source_file", "VARCHAR"), ("source_folder", "VARCHAR"),
    ("invoice_date", "DATE"), ("year", "INTEGER"), ("month", "INTEGER"),
    ("invoice_no", "VARCHAR"), ("supplier", "VARCHAR"), ("consignee", "VARCHAR"),
    ("description_raw", "VARCHAR"), ("expense_item", "VARCHAR"),
    ("expense_category", "VARCHAR"), ("amount", "DOUBLE"), ("vat", "DOUBLE"),
    ("currency", "VARCHAR"), ("order_no", "VARCHAR"), ("arrive_date", "DATE"), ("onboard_date", "DATE"),
    ("bl_no", "VARCHAR"), ("mbl_no", "VARCHAR"), ("port_of_loading", "VARCHAR"),
    ("port_of_discharge", "VARCHAR"), ("volume", "VARCHAR"), ("vessel", "VARCHAR"),
    ("container_no", "VARCHAR"), ("shipper", "VARCHAR"),
)
_DATE_COLUMNS = {name for name, typ in ITEM_COLUMNS if typ == "DATE"}

try:
    from python_calamine import CalamineWorkbook
except ImportError:  # 沒裝 calamine 時退回 openpyxl (較慢但結果相同)
    CalamineWorkbook = None


def get_logger() -> logging.Logger:
    logger = logging.getLogger("expense_etl")
    if logger.handlers:
        return logger
    logger.setLevel(logging.INFO)
    try:
        LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
        handler: logging.Handler = RotatingFileHandler(
            LOG_PATH, maxBytes=1_000_000, backupCount=5, encoding="utf-8")
    except OSError:
        handler = logging.StreamHandler()
    handler.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
    logger.addHandler(handler)
    return logger


# ============================================================
# 掃描檔案
# ============================================================

def _result_name(file_name: str) -> str:
    """原始帳單檔名 -> 轉檔工具產生的轉檔結果檔名。"""
    return f"{os.path.splitext(file_name)[0]}{RESULT_SUFFIX}"


def _is_target_folder(name: str, level: str, supplier_folder: Optional[str]) -> bool:
    if not name.endswith(f"-辨識等級({level})"):
        return False
    return supplier_folder is None or name.upper().startswith(supplier_folder.upper())


def scan_excel_files(roots: List[Path], level: str,
                     supplier_folder: Optional[str]) -> Iterator[Tuple[str, str, float, int]]:
    """逐一產生 (檔案路徑, 來源資料夾名稱, 修改時間, 大小)。用 os.scandir
    走訪，不一次把整個資料夾樹讀進記憶體。"""
    stack = [str(r) for r in roots]
    while stack:
        current = stack.pop()
        try:
            entries = list(os.scandir(current))
        except OSError as e:
            get_logger().warning("無法讀取資料夾 %s：%s", current, e)
            continue
        target = _is_target_folder(os.path.basename(current), level, supplier_folder)
        names = {e.name for e in entries}
        for entry in entries:
            if entry.is_dir(follow_symlinks=False):
                stack.append(entry.path)
            elif target and entry.name.lower().endswith((".xlsx", ".xlsm")) \
                    and not entry.name.startswith(("~$", SUMMARY_REPORT_PREFIX)) \
                    and _result_name(entry.name) not in names:
                # 最後一個條件：Excel 帳單 (越南供應商) 轉檔後，原始帳單跟
                # 「<檔名>_轉檔結果.xlsx」放在同一個資料夾，原始帳單不是轉檔結果
                st = entry.stat()
                # 來源資料夾：辨識等級資料夾的上一層 (例如客戶/專案資料夾名稱)
                folder = os.path.basename(os.path.dirname(current))
                yield entry.path, folder, st.st_mtime, st.st_size


# ============================================================
# 讀檔 + 清洗 (在子 process 裡執行)
# ============================================================

def _read_rows(path: str) -> List[list]:
    """讀第一個工作表的所有列 (值)。"""
    if CalamineWorkbook is not None:
        wb = CalamineWorkbook.from_path(path)
        return wb.get_sheet_by_index(0).to_python(skip_empty_area=False)
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _header_code(value) -> Optional[str]:
    if value is None:
        return None
    first_line = str(value).strip().split("\n")[0].strip().upper().rstrip(".").strip()
    return HEADER_FIELDS.get(first_line)


def _find_header(rows: List[list]) -> Tuple[int, Dict[int, str]]:
    """回傳 (標題列 index, {欄 index: 欄位代碼})；找不到丟出 ValueError。"""
    for i, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        mapping = {}
        for col, value in enumerate(row):
            code = _header_code(value)
            if code and code not in mapping.values():
                mapping[col] = code
        if {"description", "amount"} <= set(mapping.values()) and len(mapping) >= 4:
            return i, mapping
    raise ValueError("找不到標題列 (前幾列沒有 Description / AMOUNT 等欄位)")


def clean_record(raw: Dict[str, object],
                 default_currency: Optional[str] = DEFAULT_CURRENCY) -> Dict[str, object]:
    """一列原始值 -> 清洗後的明細 (不含 source_file/source_folder)。"""
    invoice_date = parse_date(raw.get("invoice_date"))
    if invoice_date is None and is_no_own_invoice(raw.get("supplier")):
        invoice_date = parse_date(raw.get("onboard_date"))  # 帳單沒有發票日期
    item, category = normalize_expense(raw.get("description"))
    return {
        "invoice_date": invoice_date,
        "year": invoice_date.year if invoice_date else None,
        "month": invoice_date.month if invoice_date else None,
        "invoice_no": clean_text(raw.get("invoice_no")),
        "supplier": clean_text(raw.get("supplier")),
        "consignee": clean_text(raw.get("consignee")),
        "description_raw": clean_text(raw.get("description")),
        "expense_item": item,
        "expense_category": category,
        "amount": parse_amount(raw.get("amount")),
        "vat": parse_amount(raw.get("vat")),
        "currency": detect_currency(raw.get("amount"), raw.get("supplier"), default_currency,
                                    description=raw.get("description")),
        "order_no": clean_text(raw.get("order_no")),
        "arrive_date": parse_date(raw.get("arrive_date")),
        "onboard_date": parse_date(raw.get("onboard_date")),
        "bl_no": clean_text(raw.get("bl_no")),
        "mbl_no": clean_text(raw.get("mbl_no")),
        "port_of_loading": normalize_port(raw.get("port_of_loading")),
        "port_of_discharge": normalize_port(raw.get("port_of_discharge")),
        "volume": clean_text(raw.get("volume")),  # 原文；pivot_export 用 parse_volume() 拆數字
        "vessel": clean_text(raw.get("vessel")),
        "container_no": clean_text(raw.get("container_no")),
        "shipper": clean_text(raw.get("shipper")),
    }


def load_file(task: Tuple[str, str, str]) -> Tuple[str, List[tuple], Optional[str]]:
    """讀一份 Excel 並清洗，回傳 (路徑, 明細列 tuple 清單, 錯誤原因)。
    任何例外都轉成錯誤原因回傳，不讓單一壞檔中斷整批。"""
    path, folder, default_currency = task
    # 帳單和供應商都看不出幣別時，用國家資料夾的幣別 (供應商轉檔結果\越南\... -> VND)
    default_currency = folder_currency(path) or default_currency
    try:
        rows = _read_rows(path)
        header_index, mapping = _find_header(rows)
        records = []
        for row in rows[header_index + 1:]:
            raw = {code: row[col] for col, code in mapping.items() if col < len(row)}
            # 範本後面的空白列、或只剩框線的列直接略過
            if all(clean_text(v) is None for v in raw.values()):
                continue
            rec = clean_record(raw, default_currency)
            rec["source_file"] = path
            rec["source_folder"] = folder
            records.append(tuple(rec[name] for name, _ in ITEM_COLUMNS))
        return path, records, None
    except Exception as e:  # noqa: BLE001  壞檔可能丟出各種例外
        return path, [], f"{type(e).__name__}: {e}"


# ============================================================
# 資料庫
# ============================================================

def init_db(con: duckdb.DuckDBPyConnection) -> None:
    cols = ", ".join(f"{name} {typ}" for name, typ in ITEM_COLUMNS)
    con.execute(f"CREATE TABLE IF NOT EXISTS expense_items ({cols}, "
                "is_duplicate BOOLEAN DEFAULT FALSE)")
    # 舊版資料庫沒有幣別欄：補上並填印尼盾 (舊版只有印尼盾資料)
    has_currency = con.execute(
        "SELECT count(*) FROM information_schema.columns "
        "WHERE table_name = 'expense_items' AND column_name = 'currency'").fetchone()[0]
    if not has_currency:
        con.execute("ALTER TABLE expense_items ADD COLUMN currency VARCHAR")
        con.execute("UPDATE expense_items SET currency = 'IDR'")
    con.execute("""CREATE TABLE IF NOT EXISTS etl_files (
        source_file VARCHAR PRIMARY KEY, mtime DOUBLE, size BIGINT,
        row_count INTEGER, error VARCHAR, processed_at TIMESTAMP)""")
    fx.ensure_table(con)


def _delete_files(con: duckdb.DuckDBPyConnection, paths: List[str]) -> None:
    """刪除這些檔案的明細與處理紀錄 (檔案有修改要重讀、或已不存在)。"""
    if not paths:
        return
    con.register("drop_files", pd.DataFrame({"source_file": paths}))
    con.execute("DELETE FROM expense_items WHERE source_file IN "
                "(SELECT source_file FROM drop_files)")
    con.execute("DELETE FROM etl_files WHERE source_file IN "
                "(SELECT source_file FROM drop_files)")
    con.unregister("drop_files")


def _flush(con: duckdb.DuckDBPyConnection, items: List[tuple], files: List[tuple]) -> None:
    """把累積的明細與檔案紀錄寫進資料庫，寫完清空緩衝。"""
    if items:
        df = pd.DataFrame(items, columns=[name for name, _ in ITEM_COLUMNS])
        for col in _DATE_COLUMNS:
            df[col] = pd.to_datetime(df[col])
        con.register("buf_items", df)
        select = ", ".join(f"CAST({name} AS {typ})" for name, typ in ITEM_COLUMNS)
        con.execute(f"INSERT INTO expense_items ({', '.join(n for n, _ in ITEM_COLUMNS)}) "
                    f"SELECT {select} FROM buf_items")
        con.unregister("buf_items")
        items.clear()
    if files:
        con.register("buf_files", pd.DataFrame(
            files, columns=["source_file", "mtime", "size", "row_count", "error", "processed_at"]))
        con.execute("INSERT OR REPLACE INTO etl_files SELECT * FROM buf_files")
        con.unregister("buf_files")
        files.clear()


def mark_duplicates(con: duckdb.DuckDBPyConnection) -> int:
    """同一個 (供應商, 發票號碼) 出現在好幾份 Excel 時，只保留路徑排序第一
    份，其餘的明細標記 is_duplicate。SUPPLIER_NO_OWN_INVOICE 的供應商改成整列
    (發票號碼+費用名稱+金額+提單號碼) 相同才算重複。回傳被標記的列數。"""
    con.execute("UPDATE expense_items SET is_duplicate = FALSE WHERE is_duplicate")
    row_level = list(SUPPLIER_NO_OWN_INVOICE)
    # 開頭比對 (名稱後面可能有英文備註)
    match = " OR ".join("starts_with(upper(coalesce(supplier, '')), ?)" for _ in row_level) or "FALSE"
    con.execute(f"""
        UPDATE expense_items AS e SET is_duplicate = TRUE
        FROM (
            SELECT coalesce(supplier, '') AS sup, invoice_no, min(source_file) AS keep
            FROM expense_items
            WHERE invoice_no IS NOT NULL AND NOT ({match})
            GROUP BY 1, 2
            HAVING count(DISTINCT source_file) > 1
        ) AS d
        WHERE coalesce(e.supplier, '') = d.sup AND e.invoice_no = d.invoice_no
          AND e.source_file <> d.keep""", row_level)
    if row_level:
        con.execute(f"""
            UPDATE expense_items AS e SET is_duplicate = TRUE
            FROM (
                SELECT rowid AS rid, source_file, min(source_file) OVER (
                    PARTITION BY supplier, invoice_no, description_raw, amount, bl_no) AS keep
                FROM expense_items
                WHERE invoice_no IS NOT NULL AND ({match})
            ) AS d
            WHERE e.rowid = d.rid AND d.source_file <> d.keep""", row_level)
    return con.execute("SELECT count(*) FROM expense_items WHERE is_duplicate").fetchone()[0]


def _drop_db_keep_manual_rates(db_path: Path) -> List[tuple]:
    """完整重建用：記下現有匯率 (手動、作帳匯率) 後刪掉整個資料庫檔。DuckDB 刪除資料
    不會縮小檔案，直接開新檔才不會越用越肥 (資料庫要放共用資料夾給大家讀)。"""
    if not db_path.exists():
        return []
    con = duckdb.connect(str(db_path))
    try:
        fx.ensure_table(con)
        rows = con.execute("SELECT currency, year, month, per_usd, source FROM fx_rates").fetchall()
    finally:
        con.close()
    for path in (db_path, db_path.with_name(db_path.name + ".wal")):
        if path.exists():
            path.unlink()
    return rows


# ============================================================
# 主流程
# ============================================================

def run(roots: List[Path], db_path: Path, level: str = "高",
        supplier_folder: Optional[str] = None, workers: int = 0, full: bool = False,
        parquet: Optional[Path] = None, default_currency: Optional[str] = DEFAULT_CURRENCY,
        update_fx: bool = True, use_threads: bool = False,
        echo: Callable[[str], None] = print,
        progress: Optional[Callable[[int, int], None]] = None) -> Dict:
    """彙整主流程，回傳結果摘要 dict。
    use_threads：用執行緒代替多個 process 讀檔 (儀表板/EXE 裡執行時用，避免
    在網頁伺服器或打包後的 EXE 裡另開 process)。
    echo：輸出訊息的函式；progress(已處理份數, 總份數)：進度回報。"""
    log = get_logger()
    started = time.time()
    workers = workers or max((os.cpu_count() or 2) - 1, 1)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    manual_rates = _drop_db_keep_manual_rates(db_path) if full else []
    con = duckdb.connect(str(db_path))
    try:
        init_db(con)
        if manual_rates:
            con.executemany("INSERT INTO fx_rates (currency, year, month, per_usd, source) "
                            "VALUES (?, ?, ?, ?, ?)", manual_rates)

        echo(f"掃描資料夾 (辨識等級({level})"
             + (f"、供應商資料夾 {supplier_folder}" if supplier_folder else "") + ") ...")
        scanned = {path: (folder, mtime, size) for path, folder, mtime, size
                   in scan_excel_files(roots, level, supplier_folder)}
        known = {path: (mtime, size) for path, mtime, size in
                 con.execute("SELECT source_file, mtime, size FROM etl_files").fetchall()}

        todo = [p for p, (_, mtime, size) in scanned.items() if known.get(p) != (mtime, size)]
        # 資料庫有紀錄但檔案已不存在 (被搬走/刪除) 的 -> 移除；只是這次沒掃到
        # (例如換了 --supplier-folder) 但檔案還在的保留
        removed = [p for p in known if p not in scanned and not os.path.exists(p)]
        _delete_files(con, [p for p in todo if p in known] + removed)
        echo(f"找到 {len(scanned):,} 份 Excel：新增/修改 {len(todo):,} 份、"
             f"未變動 {len(scanned) - len(todo):,} 份、已移除 {len(removed):,} 份")

        items: List[tuple] = []
        files: List[tuple] = []
        total_rows = failed = 0
        tasks = [(p, scanned[p][0], default_currency) for p in todo]
        if tasks:
            executor = ThreadPoolExecutor if use_threads else ProcessPoolExecutor
            with executor(max_workers=workers) as pool:
                for done, (path, records, error) in enumerate(
                        pool.map(load_file, tasks, chunksize=16), start=1):
                    _, mtime, size = scanned[path]
                    if error:
                        failed += 1
                        log.error("%s | 讀取失敗 | %s", path, error)
                    elif not records:
                        log.warning("%s | 沒有任何明細列", path)
                    items.extend(records)
                    total_rows += len(records)
                    files.append((path, mtime, size, len(records), error, datetime.now()))
                    if len(items) >= BATCH_ROWS:
                        _flush(con, items, files)
                    if progress:
                        progress(done, len(tasks))
                    if done % 500 == 0 or done == len(tasks):
                        rate = done / max(time.time() - started, 1e-6)
                        echo(f"  已處理 {done:,}/{len(tasks):,} 份 ({rate:,.0f} 份/秒)，"
                             f"明細 {total_rows:,} 列，失敗 {failed:,} 份")
            _flush(con, items, files)

        duplicates = mark_duplicates(con)

        fx_message = None
        if update_fx:
            try:
                fx_message = fx.update_book_rates(con, db_path.parent)
            except Exception as e:  # noqa: BLE001  匯率檔損壞等狀況不影響資料匯入
                fx_message = f"匯率更新失敗 (沿用資料庫裡現有的匯率)：{e}"
                log.warning("匯率更新失敗 | %s", e)
            echo(fx_message)

        if parquet:
            parquet.parent.mkdir(parents=True, exist_ok=True)
            con.execute(f"COPY (SELECT * FROM expense_items) TO '{parquet.as_posix()}' "
                        "(FORMAT PARQUET, COMPRESSION ZSTD)")
            echo(f"已匯出 Parquet：{parquet}")

        summary = con.execute("""
            SELECT coalesce(CAST(year AS VARCHAR), '日期無法辨識'), count(*),
                   count(DISTINCT coalesce(supplier, '') || '|' || coalesce(invoice_no, source_file)),
                   sum(amount), string_agg(DISTINCT currency, '/')
            FROM expense_items WHERE NOT is_duplicate GROUP BY 1 ORDER BY 1""").fetchall()
    finally:
        con.close()

    elapsed = time.time() - started
    echo("=" * 60)
    echo(f"完成，耗時 {elapsed:,.1f} 秒；資料庫：{db_path}")
    echo(f"本次寫入 {total_rows:,} 列，讀取失敗 {failed:,} 份 (原因見 {LOG_PATH})；"
         f"重複發票 {duplicates:,} 列不計入")
    for year, rows, invoices, amount, currencies in summary:
        echo(f"  {year}：發票 {invoices:,} 張、明細 {rows:,} 列、金額 {amount or 0:,.0f} ({currencies})")
    echo("=" * 60)
    log.info("ETL 完成 | 新增/修改 %d 份 | 寫入 %d 列 | 失敗 %d 份", len(todo), total_rows, failed)
    return {"scanned": len(scanned), "processed": len(todo), "removed": len(removed),
            "rows": total_rows, "failed": failed, "duplicates": duplicates,
            "fx": fx_message, "seconds": elapsed}


def main(argv: Optional[List[str]] = None) -> int:
    default_db = get_db_path()
    parser = argparse.ArgumentParser(description="彙整轉檔結果 Excel 進 DuckDB")
    parser.add_argument("roots", nargs="+", type=Path, help="要掃描的資料夾 (可多個，含所有子資料夾)")
    parser.add_argument("--db", type=Path, default=default_db, help=f"DuckDB 檔案 (預設 {default_db})")
    parser.add_argument("--level", default="高", help="只讀這個辨識等級的資料夾 (預設 高)")
    parser.add_argument("--supplier-folder", help="只讀名稱以此開頭的辨識等級資料夾，例如 MAERSK")
    parser.add_argument("--workers", type=int, default=0,
                        help="平行讀檔的 process 數 (預設 CPU 核心數 - 1)")
    parser.add_argument("--full", action="store_true", help="清空資料庫重新匯入")
    parser.add_argument("--parquet", type=Path, help="另外匯出一份 Parquet 檔")
    parser.add_argument("--default-currency", default=DEFAULT_CURRENCY,
                        help="帳單、供應商、國家資料夾都判斷不出幣別時用的幣別 "
                             "(預設不指定：幣別留空，再詢問使用者)")
    parser.add_argument("--no-fx", action="store_true", help="不要用作帳匯率檔更新匯率")
    args = parser.parse_args(argv)

    missing = [str(r) for r in args.roots if not r.is_dir()]
    if missing:
        print(f"找不到資料夾：{', '.join(missing)}")
        return 1
    run([r.resolve() for r in args.roots], args.db.resolve(), args.level,
        args.supplier_folder, args.workers, args.full, args.parquet,
        args.default_currency.upper() if args.default_currency else None, not args.no_fx)
    return 0


if __name__ == "__main__":
    multiprocessing.freeze_support()
    sys.exit(main())
