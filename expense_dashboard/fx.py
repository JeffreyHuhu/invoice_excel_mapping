# -*- coding: utf-8 -*-
"""
fx.py — 匯率表 (換算成 USD 用)

fx_rates 資料表：每個幣別一筆「1 USD = 多少該幣別」(per_usd)。
  - 作帳匯率 (使用者 2026-10-28 指定，唯一的自動來源)：資料庫旁邊的
          「作帳匯率 YYYYMMDD-MMDD.xlsm」，檔名日期最新的那份。讀「ALL」工作表，
          每個幣別只取最新一期，所有發票都用這個匯率換算 (不分發票月份)。
          不再上網抓匯率 (台灣銀行、ECB 都已移除)。
  - 手動：source = '手動' 的列 (在儀表板「資料管理」頁修改的)。只用在
          作帳匯率檔沒有的幣別；作帳匯率檔有的幣別一律以作帳匯率為準。
換算時找不到當月匯率就用該幣別最近一期的匯率，所以每幣別一筆就會套用到全部月份。
"""

import re
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional, Tuple

import duckdb

SOURCE_MANUAL = "手動"
SOURCE_BOOK = "作帳匯率"
BOOK_FILE_GLOB = "作帳匯率 *.xls*"
BOOK_SHEET = "ALL"
# 期間文字，例如 "2026/10/21-2026/10/31"
_PERIOD_RE = re.compile(r"(\d{4})/(\d{1,2})/(\d{1,2})\s*-\s*(\d{4})/(\d{1,2})/(\d{1,2})")
_FILE_DATE_RE = re.compile(r"(\d{8})")


def ensure_table(con: duckdb.DuckDBPyConnection) -> None:
    con.execute("""CREATE TABLE IF NOT EXISTS fx_rates (
        currency VARCHAR, year INTEGER, month INTEGER, per_usd DOUBLE,
        source VARCHAR, updated_at TIMESTAMP DEFAULT current_timestamp,
        PRIMARY KEY (currency, year, month))""")


def latest_book_file(folder: Path) -> Optional[Path]:
    """資料夾裡檔名日期最新的作帳匯率檔 (同日期再看修改時間)；沒有回傳 None。"""
    files = [f for f in Path(folder).glob(BOOK_FILE_GLOB) if not f.name.startswith("~$")]
    if not files:
        return None

    def key(f: Path):
        m = _FILE_DATE_RE.search(f.name)
        return (m.group(1) if m else "", f.stat().st_mtime)
    return max(files, key=key)


def read_book_rates(path: Path) -> Dict[str, Tuple[str, datetime, float]]:
    """作帳匯率檔 -> {幣別: (期間文字, 期間起日, 1 USD 兌換數)}，每幣別只留最新一期。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[BOOK_SHEET]
        latest: Dict[str, Tuple[str, datetime, float]] = {}
        for row in ws.iter_rows(min_row=3, max_col=4, values_only=True):
            period, src, cur, rate = (list(row) + [None] * 4)[:4]
            m = _PERIOD_RE.search(str(period or ""))
            if not m or str(src).strip().upper() != "USD" or not cur or not rate:
                continue
            begin = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            cur = str(cur).strip().upper()
            if cur not in latest or begin > latest[cur][1]:
                latest[cur] = (m.group(0), begin, float(rate))
        return latest
    finally:
        wb.close()


def update_book_rates(con: duckdb.DuckDBPyConnection, folder: Path) -> str:
    """用最新的作帳匯率檔更新 fx_rates，回傳結果說明。檔案裡有的幣別先刪掉
    全部舊匯率 (含手動)，再寫入最新一期一筆；舊的上網來源 (ECB/台灣銀行) 一併清掉。"""
    ensure_table(con)
    path = latest_book_file(folder)
    if path is None:
        return f"找不到作帳匯率檔 ({Path(folder) / BOOK_FILE_GLOB})，沿用資料庫裡現有的匯率"
    rates = read_book_rates(path)
    if not rates:
        return f"作帳匯率檔沒有可用的匯率：{path.name}，沿用資料庫裡現有的匯率"
    con.execute("BEGIN")
    con.execute(f"DELETE FROM fx_rates WHERE source <> '{SOURCE_MANUAL}'")
    con.executemany("DELETE FROM fx_rates WHERE currency = ?", [[c] for c in rates])
    con.executemany(
        "INSERT INTO fx_rates (currency, year, month, per_usd, source) VALUES (?, ?, ?, ?, ?)",
        [(c, b.year, b.month, r, f"{SOURCE_BOOK} {p}") for c, (p, b, r) in rates.items()])
    con.execute("COMMIT")
    msg = f"已套用作帳匯率 {path.name}：" + "、".join(
        f"{c} {r:g}" for c, (_, _, r) in sorted(rates.items()))
    used = [c for (c,) in con.execute(
        "SELECT DISTINCT currency FROM expense_items WHERE currency IS NOT NULL "
        "AND currency <> 'USD'").fetchall()] if con.execute(
        "SELECT count(*) FROM information_schema.tables WHERE table_name = 'expense_items'"
    ).fetchone()[0] else []
    has = {c for (c,) in con.execute("SELECT DISTINCT currency FROM fx_rates").fetchall()}
    missing = sorted(set(used) - has)
    if missing:
        msg += f"；作帳匯率檔沒有 {', '.join(missing)}，請在儀表板「資料管理」頁手動輸入"
    return msg


def rate_sql(fixed: Dict[str, float] = None) -> Tuple[str, list]:
    """換算成 USD 的 items CTE：amount / vat 換成 USD，另加 fx_rate (1 USD =
    多少原幣) 與 amount_orig (原幣金額)。fixed：{幣別: 1 USD 兌換數}，有
    給的幣別一律用固定匯率。回傳 (SQL, 參數)。"""
    fixed = fixed or {}
    fixed_cases = " ".join("WHEN e.currency = ? THEN ?" for _ in fixed)
    params: list = []
    for cur, rate in fixed.items():
        params.extend([cur, rate])
    sql = f"""
    latest_fx AS (
        SELECT currency, arg_max(per_usd, year * 100 + month) AS per_usd
        FROM fx_rates GROUP BY currency),
    items_fx AS (
        SELECT e.*, CASE WHEN e.currency = 'USD' THEN 1.0 {fixed_cases}
                         ELSE coalesce(f.per_usd, l.per_usd) END AS fx_rate
        FROM expense_items e
        LEFT JOIN fx_rates f ON f.currency = e.currency AND f.year = e.year AND f.month = e.month
        LEFT JOIN latest_fx l ON l.currency = e.currency),
    items AS (
        SELECT * EXCLUDE (amount, vat), amount AS amount_orig,
               amount / fx_rate AS amount, vat / fx_rate AS vat
        FROM items_fx)"""
    return sql, params
