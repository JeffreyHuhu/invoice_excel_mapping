# -*- coding: utf-8 -*-
"""
combine_results.py — 把所有轉檔結果 Excel 整合成一份 (欄位跟轉檔結果模板一樣)

【做什麼】
  1. 掃描資料夾底下「XXX-辨識等級(高)」「XXX-辨識等級(中)」裡的轉檔結果 Excel
     (沿用 process_data.scan_excel_files，規則一樣：跳過轉檔總表、原始 Excel 帳單)。
  2. 依標題列找出 18 個標準欄位，順序與抬頭跟轉檔結果模板一樣；最右邊加一欄
     「辨識等級」(高/中)，讓主管分辨可能有錯的資料 (使用者 2026-10-28 指定)。
  3. 同一張發票 (供應商 + 發票號碼) 出現在多份 Excel 時，只保留第一份
     (等級高優先，再依路徑排序)，跟資料庫的重複判斷一樣 (含逐列比對的供應商)。
  4. 寫成一個工作表「工作表」，日期欄位用 mmm-dd-yyyy 顯示。
     讀不了的檔案記到 log，不中斷整批處理。
  5. 有給 totals_fx_folder (個別供應商的「<KEY>轉檔結果整合_資料模板.xlsx」) 時加第二個
     工作表「費用總和」：依幣別列出費用 (AMOUNT)、VAT、TOTAL (= 費用 + VAT) 的
     總和，當地幣別與美金各一組 (使用者 2026-10-28 指定)。匯率用作帳匯率檔
     (fx.latest_book_file / read_book_rates，跟資料庫一樣，每幣別最新一期)。

【怎麼用】
  python combine_results.py <輸出.xlsx> [<資料夾1> <資料夾2> ...] [--levels 高 中]
"""

import argparse
import sys
from datetime import date, datetime
from pathlib import Path
from typing import List, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font

from expense_normalize import detect_currency, folder_currency, is_no_own_invoice, parse_amount
from process_data import _find_header, _read_rows, get_logger, scan_excel_files

# 抬頭跟 extract_utils_streamlit.DISPLAY_HEADERS 一致 (那支檔案要載入 PDF 套件，這裡不 import)
TEMPLATE_HEADERS = {
    "invoice_date": "INVOICE DATE\n發票日期",
    "invoice_no": "INVOICE NO\n發票號碼",
    "supplier": "SUPPLIER\n供應商名稱",
    "consignee": "CONSIGNEE\n集團公司名稱",
    "description": "Description\n費用名稱",
    "amount": "AMOUNT\n金額",
    "vat": "VAT\n稅額",
    "order_no": "Order No.",
    "arrive_date": "Arrive Date\n到達日",
    "onboard_date": "on board date\n開船日",
    "bl_no": "B/L NO.\n提單號碼",
    "mbl_no": "MBL NO.\n主提單號碼",
    "port_of_loading": "Port of Loading\n啟運港",
    "port_of_discharge": "Port of discharge\n目的港",
    "volume": "Volume\n材積",
    "vessel": "Vessel\n船名",
    "container_no": "Container no.\n貨櫃號碼",
    "shipper": "Shipper\n出貨方",
}
FIELDS = list(TEMPLATE_HEADERS)
LEVEL_HEADER = "辨識等級"
DATE_FIELDS = {"invoice_date", "arrive_date", "onboard_date"}
DATE_FORMAT = "mmm-dd-yyyy"
SHEET_NAME = "工作表"
TOTALS_SHEET_NAME = "費用總和"
EXCEL_MAX_ROWS = 1_048_575
FONT_NAME = "新細明體"
SHEET_WIDTHS = {"A": 18.0, "B": 12.7, "C": 14.5, "G": 15.7, "I": 12.3, "J": 17.2, "K": 16.5,
                "L": 19.3, "M": 14.8, "N": 18.7, "Q": 15.3}


def _invoice_key(row: dict):
    """(供應商, 發票號碼)；沒有發票號碼就不做重複判斷。SUPPLIER_NO_OWN_INVOICE
    的供應商要整列 (發票號碼+費用名稱+金額+提單號碼) 相同才算重複 (跟資料庫一樣)。"""
    no = row.get("invoice_no")
    if no in (None, ""):
        return None
    key = (str(row.get("supplier") or "").strip().upper(), str(no).strip().upper())
    if is_no_own_invoice(row.get("supplier")):
        key += tuple(str(row.get(f)) for f in ("description", "amount", "bl_no"))
    return key


def collect_rows(roots: Sequence[Path], levels: Sequence[str]):
    """回傳 (明細列, 統計)。明細列是 {欄位代碼: 值}。"""
    log = get_logger()
    files = []
    for rank, level in enumerate(levels):
        files += [(rank, level, path) for path, _, _, _ in scan_excel_files(list(roots), level, None)]
    files.sort(key=lambda t: (t[0], t[2]))

    rows: List[dict] = []
    seen = set()
    stats = {lv: {"files": 0, "rows": 0} for lv in levels}
    stats.update(failed=0, dup_rows=0)
    for _, level, path in files:
        try:
            raw = _read_rows(path)
            header_idx, mapping = _find_header(raw)
        except Exception as e:  # 損壞或格式不符：記 log，繼續下一份
            log.warning("整合 Excel 讀取失敗 %s：%s", path, e)
            stats["failed"] += 1
            continue
        file_rows = []
        for values in raw[header_idx + 1:]:
            row = {code: values[col] for col, code in mapping.items() if col < len(values)}
            row = {k: (None if v == "" else v) for k, v in row.items()}
            if any(v is not None for v in row.values()):
                row["level"] = level
                row["_path"] = path          # 判斷幣別用 (國家資料夾)，不寫進 Excel
                file_rows.append(row)
        keys = {_invoice_key(r) for r in file_rows} - {None}
        dup_keys = keys & seen
        if dup_keys:
            before = len(file_rows)
            file_rows = [r for r in file_rows if _invoice_key(r) not in dup_keys]
            stats["dup_rows"] += before - len(file_rows)
        seen |= keys
        rows += file_rows
        stats[level]["files"] += 1
        stats[level]["rows"] += len(file_rows)
    return rows, stats


def _ticket_key(row: dict):
    """票數的單位：B/L NO. (有數字的才算)；沒有就用 INVOICE NO
    (跟本機轉檔總表 local_quality._shipment_count 一樣)。"""
    bl = str(row.get("bl_no") or "").strip()
    if bl and bl.upper() != "N/A" and any(ch.isdigit() for ch in bl):
        return ("BL", bl.upper())
    inv = str(row.get("invoice_no") or "").strip()
    return ("INV", inv.upper()) if inv and inv.upper() != "N/A" else None


def compute_totals(rows: List[dict], fx_folder: Path):
    """依幣別加總 -> ([(幣別, 總票數 (不重複), 費用, VAT, TOTAL, 1 USD 兌換數或 None)], 匯率檔名)。
    幣別判斷跟資料庫一樣：帳單幣別 > 供應商 > 國家資料夾。"""
    from collections import OrderedDict
    import fx
    sums = OrderedDict()
    for row in rows:
        cur = detect_currency(row.get("amount"), row.get("supplier"),
                              folder_currency(row.get("_path")),
                              description=row.get("description")) or "(未知)"
        s = sums.setdefault(cur, [set(), 0.0, 0.0])
        key = _ticket_key(row)
        if key:
            s[0].add(key)
        s[1] += parse_amount(row.get("amount")) or 0.0
        s[2] += parse_amount(row.get("vat")) or 0.0
    book = fx.latest_book_file(fx_folder)
    rates = fx.read_book_rates(book) if book else {}
    out = []
    for cur, (tickets, amount, vat) in sums.items():
        rate = 1.0 if cur == "USD" else (rates[cur][2] if cur in rates else None)
        out.append((cur, len(tickets), amount, vat, amount + vat, rate))
    return out, (book.name if book else None)


def supplier_label(key: str) -> str:
    """費用總和標題用的「VN-05 CONG_THANH」：編號取 suppliers(_streamlit).py 的
    SUPPLIER_NUMBERS (用 ast 讀，不載入 PDF 套件)；找不到編號就只用 KEY。"""
    import ast
    import re
    for name in ("suppliers_streamlit.py", "suppliers.py"):
        path = Path(__file__).resolve().parent.parent / name
        try:
            for node in ast.parse(path.read_text(encoding="utf-8")).body:
                if (isinstance(node, ast.Assign) and getattr(node.targets[0], "id", "") == "SUPPLIER_NUMBERS"):
                    number = ast.literal_eval(node.value).get(key, ("",))[0]
                    number = re.sub(r"^([A-Z]+)(\d)", r"\1-\2", number)
                    return f"{number} {key}".strip()
        except (OSError, SyntaxError, ValueError):
            continue
    return key


def _write_totals_sheet(wb: Workbook, totals, book_name, supplier_key: str = None) -> None:
    """格式照使用者 2026-10-29 改好的 CONG_THANH 範例：黃底標題「VN-05 CONG_THANH 總和
    (當地幣別 & 美金)」；三個區塊 (總票數/幣別、當地幣別金額、USD 金額) 各一種淺色底
    (主題色 9/6/8，淡化 80%)，新細明體、全部靠右，區塊之間空兩列。"""
    from openpyxl.styles import Color, PatternFill
    ws = wb.create_sheet(TOTALS_SHEET_NAME)
    for col, width in zip("ABCDEFGHI", (20.1, 19.7, 28.8, 24.0, 23.3, 20.0, 20.4, 21.0, 23.9)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 19.5
    font_name = FONT_NAME
    yellow = PatternFill("solid", fgColor="FFFFFF00")
    fills = [PatternFill("solid", fgColor=Color(theme=t, tint=0.7999816888943144)) for t in (9, 6, 8)]
    right = Alignment(horizontal="right")

    def cell(value, fill, bold=False, fmt=None, size=11, align=right):
        c = WriteOnlyCell(ws, value=value)
        c.font = Font(name=font_name, bold=bold, size=size)
        c.fill = fill
        if align is not None:
            c.alignment = align
        if fmt:
            c.number_format = fmt
        return c

    def block(fill, header, values, fmt=None):
        ws.append([cell(h, fill, bold=True) for h in header])
        for vals in values:
            ws.append([cell(v, fill, fmt=fmt if isinstance(v, (int, float)) and fmt else None)
                       for v in vals])

    money = "#,##0.00"
    title = f"{supplier_label(supplier_key)} 總和" if supplier_key else "費用總和"
    ws.append([cell(f"{title} (當地幣別 & 美金)", yellow, bold=True, size=14, align=None),
               cell(None, yellow, align=None), cell(None, yellow, align=None)])
    ws.append([])
    block(fills[0], ["總票數 (不重複)", "幣別"], [[n, cur] for cur, n, *_ in totals])
    missing = []
    usd_sum = [0.0, 0.0, 0.0]
    for cur, n, amount, vat, total, rate in totals:
        ws.append([])
        ws.append([])
        block(fills[1], [f"費用項目 ({cur})", f"VAT ({cur})", f"TOTAL 總和  ({cur})"],
              [[amount, vat, total]], money)
        if rate:
            usd_sum = [x + v / rate for x, v in zip(usd_sum, (amount, vat, total))]
        else:
            missing.append(cur)
    ws.append([])
    ws.append([])
    block(fills[2], ["費用項目 (USD)", "VAT (USD)", "TOTAL 總和 (USD)"], [usd_sum], money)
    if missing:
        ws.append([])
        ws.append([f"注意：{'、'.join(missing)} 沒有匯率 ({book_name or '找不到作帳匯率檔'})，沒有算進 USD"])


def write_workbook(rows: List[dict], out_path: Path, totals=None, supplier_key: str = None) -> None:
    """write_only 模式寫檔，大量資料也不佔太多記憶體。"""
    if len(rows) > EXCEL_MAX_ROWS:
        raise SystemExit(f"明細 {len(rows):,} 列，超過 Excel 上限 {EXCEL_MAX_ROWS:,} 列。")
    # 樣式照使用者 2026-10-29 改好的 CONG_THANH 範例：全部新細明體、表頭列高 58、
    # 篩選整個表、指定欄寬 (其餘欄用預設寬度)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(SHEET_NAME)
    ws.freeze_panes = "A2"
    for col, width in SHEET_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 58
    ws.auto_filter.ref = f"A1:{chr(ord('A') + len(FIELDS))}{len(rows) + 1}"
    bold = Font(name=FONT_NAME, bold=True, size=11)
    plain = Font(name=FONT_NAME, size=11)
    wrap = Alignment(wrap_text=True, vertical="center")
    header = []
    for text in [*TEMPLATE_HEADERS.values(), LEVEL_HEADER]:
        cell = WriteOnlyCell(ws, value=text)
        cell.font, cell.alignment = bold, wrap
        header.append(cell)
    ws.append(header)
    for row in rows:
        out = []
        for field in FIELDS:
            value = row.get(field)
            cell = WriteOnlyCell(ws, value=value)
            cell.font = plain
            if field in DATE_FIELDS and isinstance(value, (date, datetime)):
                cell.number_format = DATE_FORMAT
            out.append(cell)
        cell = WriteOnlyCell(ws, value=row["level"])
        cell.font = plain
        out.append(cell)
        ws.append(out)
    if totals is not None:
        _write_totals_sheet(wb, *totals, supplier_key=supplier_key)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def run(out_path: Path, roots: Sequence[Path], levels: Sequence[str] = ("高", "中"),
        totals_fx_folder: Path = None, supplier_key: str = None) -> Path:
    """totals_fx_folder：放作帳匯率檔的資料夾；有給就加「費用總和」工作表。
    supplier_key：個別供應商的 KEY，放進「費用總和」標題 (例如 VN-05 CONG_THANH)。"""
    rows, stats = collect_rows(roots, levels)
    summary = compute_totals(rows, totals_fx_folder) if totals_fx_folder else None
    write_workbook(rows, out_path, summary, supplier_key)
    print(f"已整合轉檔結果 Excel：{out_path}")
    for lv in levels:
        print(f"  等級({lv})：{stats[lv]['files']:,} 份、{stats[lv]['rows']:,} 列")
    print(f"  共 {len(rows):,} 列；重複發票略過 {stats['dup_rows']:,} 列；讀取失敗 {stats['failed']:,} 份")
    for cur, n, amount, vat, total, rate in (summary[0] if summary else []):
        usd = f"{total / rate:,.2f} USD" if rate else "沒有匯率"
        print(f"  {cur}：總票數 {n:,}、費用 {amount:,.2f}、VAT {vat:,.2f}、TOTAL {total:,.2f} ({usd})")
    return out_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("out", type=Path)
    parser.add_argument("roots", nargs="+", type=Path)
    parser.add_argument("--levels", nargs="+", default=["高", "中"])
    args = parser.parse_args()
    run(args.out, args.roots, args.levels)
    return 0


if __name__ == "__main__":
    sys.exit(main())
