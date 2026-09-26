# -*- coding: utf-8 -*-
"""
extract_utils.py
=================
多供應商帳單擷取系統：共用核心引擎 (供應商無關)

流程：辨識哪間供應商帳單 → 選取對應的帳單辨識系統 → 擷取欄位 → 跟正確答案比對算正確率

架構：
  - 這支檔案只放「跟供應商無關」的共用邏輯：PDF文字擷取、供應商辨識/註冊表、
    迴圈找最佳擷取設定、比對評分、Excel讀寫。
  - 每個供應商專屬的辨識規則 (detect) 與擷取規則 (parse) 放在 suppliers.py，
    透過 register_supplier() 註冊進本檔案的 SUPPLIER_REGISTRY。
  - 新增第 21 家、22 家...供應商時，只需要在 suppliers.py 多寫一組
    detect()/parse() 並呼叫 register_supplier()，不用改這支檔案，
    也不用改 app.py。

標準欄位 (不管哪個供應商，擷取結果都要對應同一套欄位代碼，這樣才能放進
同一份 Excel、用同一套比對/查核邏輯)：
  invoice_date, invoice_no, supplier, consignee, order_no, arrive_date,
  onboard_date, bl_no, mbl_no, port_of_loading, port_of_discharge, volume,
  vessel, container_no  → 每張帳單「共用」一次 (抬頭欄位)
  description, amount   → 每筆費用明細各一列 (一張帳單可能有多筆費用)
"""

import re
import io
import datetime
from typing import Callable, Dict, List, Optional, Tuple

import pdfplumber
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NA = "N/A"  # 抓不到值時，一律填這個 (使用者需求)

# ---------------------------------------------------------------------------
# 1. 標準欄位定義
# ---------------------------------------------------------------------------

HEADER_FIELD_CODES: List[str] = [
    "invoice_date", "invoice_no", "supplier", "consignee", "order_no",
    "arrive_date", "onboard_date", "bl_no", "mbl_no", "port_of_loading",
    "port_of_discharge", "volume", "vessel", "container_no",
]
ITEM_FIELD_CODES: List[str] = ["description", "amount"]

# 完整欄位順序 (輸出 Excel 用，對應各供應商範本欄位順序：
# INVOICE DATE, INVOICE NO, SUPPLIER, CONSIGNEE, Description, AMOUNT, Order No.,
# Arrive Date, on board date, B/L NO., MBL NO., Port of Loading,
# Port of discharge, Volume, Vessel, Container no.)
FIELD_CODES: List[str] = [
    "invoice_date", "invoice_no", "supplier", "consignee", "description",
    "amount", "order_no", "arrive_date", "onboard_date", "bl_no", "mbl_no",
    "port_of_loading", "port_of_discharge", "volume", "vessel", "container_no",
]

DISPLAY_HEADERS: Dict[str, str] = {
    "invoice_date":      "INVOICE DATE\n發票日期",
    "invoice_no":        "INVOICE NO\n發票號碼",
    "supplier":          "SUPPLIER\n供應商名稱",
    "consignee":         "CONSIGNEE\n集團公司名稱",
    "description":       "Description\n費用名稱",
    "amount":            "AMOUNT\n金額",
    "order_no":          "Order No.",
    "arrive_date":       "Arrive Date\n到達日",
    "onboard_date":      "on board date\n開船日",
    "bl_no":             "B/L NO.\n提單號碼",
    "mbl_no":            "MBL NO.\n主提單號碼",
    "port_of_loading":   "Port of Loading\n啟運港",
    "port_of_discharge": "Port of discharge\n目的港",
    "volume":            "Volume\n材積",
    "vessel":             "Vessel\n船名",
    "container_no":      "Container no.\n貨櫃號碼",
}


# ---------------------------------------------------------------------------
# 2. PDF 文字擷取 (可調參數，供迴圈嘗試不同設定用)
# ---------------------------------------------------------------------------

def extract_text_from_pdf(pdf_path: str, **pdfplumber_kwargs) -> str:
    """讀取 PDF 全部頁面文字並合併成一個字串。

    帳單是「可編輯 PDF」，pdfplumber 可以直接抓到文字層。若遇到沒有文字層
    的掃描檔，這裡會回傳空字串，使用端應提示改用 OCR 或人工輸入。

    pdfplumber_kwargs 轉給 page.extract_text()，例如 x_tolerance、
    y_tolerance、layout。不同供應商/不同份 PDF 的版面留白不盡相同，同一組
    參數不一定每份都是最佳解，因此開放參數化，讓 find_best_extraction()
    可以多試幾組，挑出效果最好的。
    """
    parts = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text(**pdfplumber_kwargs) or "")
    return "\n".join(parts)


EXTRACTION_CONFIGS: List[Dict] = [
    {},                                    # pdfplumber 預設值
    {"x_tolerance": 1, "y_tolerance": 1},
    {"x_tolerance": 1, "y_tolerance": 3},
    {"x_tolerance": 2, "y_tolerance": 3},
    {"x_tolerance": 3, "y_tolerance": 5},
    {"layout": True},                       # 盡量保留原始版面座標間距
]


# ---------------------------------------------------------------------------
# 3. 供應商辨識與註冊表
# ---------------------------------------------------------------------------
#
# 每個供應商模組 (suppliers.py) 呼叫 register_supplier() 註冊三樣東西：
#   - key        內部代碼，例如 "DWIHARTA"
#   - label      顯示名稱
#   - detect_fn  detect(text) -> bool，判斷這份PDF文字是不是這家供應商
#   - parse_fn   parse(text) -> {"header": {...}, "items": [{...}, ...]}
#                header 裡的 key 對應 HEADER_FIELD_CODES，
#                items 是一個 list，每筆是 {"description":.., "amount":..}，
#                至少要回傳一筆 (抓不到費用明細就回傳一筆全空的)。

SUPPLIER_REGISTRY: Dict[str, Dict] = {}


def register_supplier(key: str, label: str,
                       detect_fn: Callable[[str], bool],
                       parse_fn: Callable[[str], Dict]) -> None:
    SUPPLIER_REGISTRY[key] = {"label": label, "detect": detect_fn, "parse": parse_fn}


def detect_supplier(text: str) -> Optional[str]:
    """依序問過每個已註冊的供應商『這是你的帳單嗎』，回傳第一個承認的 key。
    都沒有承認就回傳 None，呼叫端應提示『無法辨識供應商』。
    """
    for key, entry in SUPPLIER_REGISTRY.items():
        try:
            if entry["detect"](text):
                return key
        except Exception:
            continue
    return None


def supplier_label(key: Optional[str]) -> str:
    if key and key in SUPPLIER_REGISTRY:
        return SUPPLIER_REGISTRY[key]["label"]
    return "未知供應商 (無法辨識)"


def list_registered_suppliers() -> List[Dict[str, str]]:
    """回傳目前已註冊的供應商清單 (給網頁上顯示用)。"""
    return [{"key": k, "label": v["label"]} for k, v in SUPPLIER_REGISTRY.items()]


# ---------------------------------------------------------------------------
# 4. 擷取結果組裝：header + items -> 完整列 (一筆費用一列)，缺漏補 N/A
# ---------------------------------------------------------------------------

def _fill_na(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return NA
    return value


def expand_to_rows(parsed: Dict) -> List[Dict[str, object]]:
    """把供應商 parse() 回傳的 {"header":.., "items":[..]} 展開成
    每筆費用一列的完整 row (欄位對應 FIELD_CODES，缺漏一律補 N/A)。
    """
    header = parsed.get("header", {}) or {}
    items = parsed.get("items") or [{}]

    header_row = {code: _fill_na(header.get(code)) for code in HEADER_FIELD_CODES}

    rows = []
    for item in items:
        row = dict(header_row)
        row["description"] = _fill_na(item.get("description"))
        row["amount"] = _fill_na(item.get("amount"))
        rows.append(row)
    return rows


def parse_invoice_pdf(pdf_path: str, supplier_key: Optional[str] = None,
                       **pdfplumber_kwargs) -> Tuple[Optional[str], List[Dict], str]:
    """讀取PDF、自動判斷供應商 (若未指定)、用對應規則擷取欄位。
    回傳 (辨識到的供應商key或None, 展開後的rows列表, 原始文字)。
    """
    text = extract_text_from_pdf(pdf_path, **pdfplumber_kwargs)
    key = supplier_key or detect_supplier(text)
    if key and key in SUPPLIER_REGISTRY:
        parsed = SUPPLIER_REGISTRY[key]["parse"](text)
    else:
        parsed = {"header": {}, "items": [{}]}
    rows = expand_to_rows(parsed)
    return key, rows, text


# ---------------------------------------------------------------------------
# 5. 比對評分：擷取結果 vs 正確答案
# ---------------------------------------------------------------------------

def normalize_value(value) -> str:
    """統一轉成好比對的字串：去除空白/大小寫差異，日期物件轉成固定格式，
    數字不管原本是 int/float/字串、有沒有多餘的小數位 (例如 '11.810' 對
    11.81、'1076510528' 對 1076510528)，都統一轉成同一種數字字串比較，
    這樣不會因為型態或多餘的零而誤判成不相符。
    """
    if value is None:
        return ""
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime("%d.%m.%Y")

    # 先嘗試當數字比對 (int/float本身，或看起來像數字的字串)
    if isinstance(value, (int, float)):
        num = float(value)
        return f"{num:.6f}".rstrip("0").rstrip(".")

    text = str(value).strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", text):
        num = float(text)
        return f"{num:.6f}".rstrip("0").rstrip(".")

    text = text.upper()
    text = re.sub(r"\s+", " ", text)
    return text.rstrip(",")


def compare_rows(rows: List[Dict], reference_rows: pd.DataFrame) -> List[Dict]:
    """比對「一張帳單」擷取出來的 rows (可能有多筆費用) 與正確答案
    (reference_rows：同一張發票在正確答案 Excel 裡的所有列)。

    抬頭欄位只比對一次 (取第一列)；費用明細 (description+amount) 用集合
    比對，不要求擷取順序跟正確答案一致。
    """
    if not rows or reference_rows is None or reference_rows.empty:
        return []

    results = []
    first_row = rows[0]
    ref_first = reference_rows.iloc[0].to_dict()

    for code in HEADER_FIELD_CODES:
        ext_val = first_row.get(code, NA)
        ref_val = ref_first.get(code, NA)
        match = normalize_value(ext_val) == normalize_value(ref_val)
        results.append({
            "欄位": DISPLAY_HEADERS[code].split("\n")[0],
            "擷取結果": ext_val,
            "正確答案": ref_val,
            "結果": "✅ 相符" if match else "❌ 不相符",
        })

    ext_items = {(normalize_value(r.get("description")), normalize_value(r.get("amount")))
                 for r in rows}
    ref_items = {(normalize_value(r.get("description")), normalize_value(r.get("amount")))
                 for _, r in reference_rows.iterrows()}

    for desc, amt in sorted(ext_items & ref_items):
        results.append({"欄位": "Description/Amount", "擷取結果": f"{desc} / {amt}",
                         "正確答案": f"{desc} / {amt}", "結果": "✅ 相符"})
    for desc, amt in sorted(ext_items - ref_items):
        results.append({"欄位": "Description/Amount", "擷取結果": f"{desc} / {amt}",
                         "正確答案": "(正確答案沒有)", "結果": "❌ 不相符"})
    for desc, amt in sorted(ref_items - ext_items):
        results.append({"欄位": "Description/Amount", "擷取結果": "(未擷取到)",
                         "正確答案": f"{desc} / {amt}", "結果": "❌ 不相符"})

    return results


def compute_accuracy(results: List[Dict]) -> float:
    if not results:
        return 0.0
    matched = sum(1 for r in results if r["結果"].startswith("✅"))
    return matched / len(results)


def quality_score(rows: List[Dict]) -> float:
    """沒有正確答案 Excel 時的自我檢查分數：欄位有抓到值(不是N/A)的比例。
    僅供迴圈挑選較好的擷取設定用，不是跟人工核對過的正確率。
    """
    if not rows:
        return 0.0
    total = ok = 0
    for row in rows:
        for v in row.values():
            total += 1
            if v and v != NA:
                ok += 1
    return ok / total if total else 0.0


# ---------------------------------------------------------------------------
# 6. 迴圈：反覆嘗試不同擷取設定，直到 100% 正確率 (或找出目前能達到的最高值)
# ---------------------------------------------------------------------------

def find_best_extraction(
    pdf_path: str,
    supplier_key: Optional[str] = None,
    reference_rows: Optional[pd.DataFrame] = None,
    configs: Optional[List[Dict]] = None,
):
    """對同一份 PDF 用多組擷取參數重複解析，每次都評分，一旦達到 100% 正確率
    就提早停止；跑完所有設定仍未達 100% 的話，回傳嘗試過最高分數的結果。

    回傳：(辨識到的supplier_key, 最佳rows, 最佳分數0~1, 使用的設定,
           逐欄比對明細或None, 全部嘗試紀錄)
    """
    configs = configs or EXTRACTION_CONFIGS
    attempts: List[Dict] = []
    best = None
    detected_key = supplier_key

    for cfg in configs:
        text = extract_text_from_pdf(pdf_path, **cfg)
        key = supplier_key or detect_supplier(text)
        detected_key = detected_key or key

        if key and key in SUPPLIER_REGISTRY:
            parsed = SUPPLIER_REGISTRY[key]["parse"](text)
        else:
            parsed = {"header": {}, "items": [{}]}
        rows = expand_to_rows(parsed)

        has_ref = reference_rows is not None and not reference_rows.empty
        if has_ref:
            results = compare_rows(rows, reference_rows)
            score = compute_accuracy(results)
        else:
            results = None
            score = quality_score(rows)

        attempts.append({"設定": cfg or "預設", "分數(%)": round(score * 100, 1)})

        if best is None or score > best["score"]:
            best = {"rows": rows, "score": score, "config": cfg, "results": results}

        if best["score"] >= 1.0:
            break  # 已經 100% 正確率，不用再嘗試其他設定

    return (detected_key, best["rows"], best["score"], best["config"],
            best["results"], attempts)


# ---------------------------------------------------------------------------
# 7. Excel 讀寫
# ---------------------------------------------------------------------------

_HEADER_ANCHOR = "INVOICE DATE"


def read_reference_excel(source) -> pd.DataFrame:
    """讀取『正確答案』Excel。不同供應商範本的表頭可能在第1列或第2列、
    欄位可能從A欄或B欄開始，這裡自動搜尋含有 'INVOICE DATE' 字樣的儲存格
    當作錨點定位，不用替每個供應商範本各寫一次死板的座標。
    """
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    wb = load_workbook(source, data_only=True)
    ws = wb.active

    header_row, start_col = None, None
    for r in range(1, 5):
        for c in range(1, 5):
            val = ws.cell(row=r, column=c).value
            if val and _HEADER_ANCHOR in str(val).upper():
                header_row, start_col = r, c
                break
        if header_row:
            break
    if header_row is None:
        header_row, start_col = 1, 1

    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
        rec = {}
        empty = True
        for idx, code in enumerate(FIELD_CODES):
            cell = row[start_col - 1 + idx]
            rec[code] = cell.value
            if cell.value not in (None, ""):
                empty = False
        if not empty:
            records.append(rec)
    return pd.DataFrame(records)


def build_excel(all_rows: List[Dict[str, object]]) -> bytes:
    """輸出成統一格式的 Excel：A欄是辨識到的供應商名稱，B欄起是標準欄位，
    表頭在第1列，資料從第2列開始。不管來源是哪個供應商都放在同一張表，
    方便一次彙整多家供應商、多份帳單的轉檔結果。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "帳單彙整"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    wrap_center = Alignment(wrap_text=True, vertical="center", horizontal="center")

    ws.cell(row=1, column=1, value="供應商").font = bold
    ws.cell(row=1, column=1).fill = header_fill
    ws.column_dimensions["A"].width = 22

    for idx, code in enumerate(FIELD_CODES, start=2):
        cell = ws.cell(row=1, column=idx, value=DISPLAY_HEADERS[code])
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = wrap_center
        ws.column_dimensions[get_column_letter(idx)].width = 20

    for r, row in enumerate(all_rows, start=2):
        ws.cell(row=r, column=1, value=row.get("__supplier_label", NA))
        for idx, code in enumerate(FIELD_CODES, start=2):
            ws.cell(row=r, column=idx, value=row.get(code, NA))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
