# -*- coding: utf-8 -*-
"""
verify_app.py
=============
供應商帳單擷取核對系統 (人工查核 + 回饋收集工具)

這是跟「供應商帳單自動化辨識系統」(app.py) 分開的獨立 Streamlit App，
專門給人工「逐欄」核對系統擷取出來的資料是否正確，用於：

  1. 上傳一份供應商帳單 PDF (不需要正確答案 Excel)。
  2. 系統自動辨識供應商、擷取欄位 (跟 app.py 用同一套 suppliers.py /
     extract_utils.py 核心引擎，共用同一套辨識/擷取規則，兩邊不會有
     擷取邏輯不一致的問題)。
  3. 人工逐欄核對：每個欄位預設「沒有動作 = 視為忽略 (不計入回饋)」，
     只有按下「❌ 標記為錯誤」才會展開輸入框，讓使用者填入正確答案。
  4. 核對完成後，把「所有被標記為錯誤的欄位 + 使用者填的正確答案」匯出
     成一份 Excel (回饋記錄)，可以把這份 Excel 拿給開發端 (或直接請
     Claude 執行 add-invoice-supplier 這套流程) 依照回饋記錄修正
     suppliers.py 裡對應供應商的 parse() 規則，改完之後回到這個網頁、
     重新上傳同一份 PDF 再核對一次，確認錯誤欄位已經修正。

跟 app.py 的差異：
  - app.py 是「有正確答案 Excel 時的自動化比對」，回答「正確率幾%」。
  - verify_app.py 是「沒有正確答案 Excel 時的人工逐欄查核」，回答「這一
    欄到底對不對、不對的話正確答案是什麼」，產出的 Excel 就是下一次可以
    拿來當作「正確答案」或「除錯依據」的原始素材。
  - 兩支 App 各自獨立部署 (Streamlit Cloud 各自建一個 App，Main file
    path 分別指向 app.py / verify_app.py)，但共用同一個 GitHub repo、
    同一套 suppliers.py / extract_utils.py，新增或修正供應商規則時，
    兩邊會同時生效，不用分別維護兩份擷取邏輯。
"""

import hashlib
import io
import tempfile

import streamlit as st

# ---------------------------------------------------------------------------
# 匯入失敗時，印出清楚的錯誤訊息 (原因/排解方式跟 app.py 完全一樣：通常是
# Streamlit Cloud 沒有重新安裝 requirements.txt，到 Manage app 點 Reboot
# app 即可)，不要讓 Streamlit 預設的模糊訊息蓋掉真正的錯誤內容。
# ---------------------------------------------------------------------------
try:
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    import suppliers  # noqa: F401 (觸發所有供應商的 register_supplier())
    from extract_utils import (
        FIELD_CODES,
        HEADER_FIELD_CODES,
        DISPLAY_HEADERS,
        NA,
        find_best_extraction,
        detect_supplier,
        extract_text_from_pdf,
        supplier_label,
        list_registered_suppliers,
        is_multi_page_supplier,
    )
except ImportError as e:
    st.set_page_config(page_title="供應商帳單擷取核對系統 - 啟動失敗", layout="wide")
    st.error(
        f"❌ 系統啟動失敗，缺少必要的 Python 套件：`{e}`\n\n"
        "這通常代表 Streamlit Cloud 沒有正確安裝 `requirements.txt` 裡列出的套件"
        "(常發生在 requirements.txt 是後來才新增/修改，但環境沒有重新完整建置時)。"
        "請到這個 App 頁面右下角點 **Manage app** → 選單裡點 **Reboot app**，"
        "強制重新安裝所有套件。"
    )
    st.stop()

st.set_page_config(page_title="供應商帳單擷取核對系統", layout="wide")

# ---------------------------------------------------------------------------
# 全站樣式：字體/按鈕顏色沿用「供應商帳單自動化辨識系統」(app.py) 同一套
# 視覺規範 (淺藍色按鈕、16px 字體)，讓兩支 App 看起來是同一套系統的兩個
# 頁面，不會有風格不一致的違和感。
# ---------------------------------------------------------------------------
st.markdown(
    """
    <style>
    div.stButton > button,
    div.stDownloadButton > button,
    div.stFormSubmitButton > button,
    div[data-testid="stButton"] > button,
    div[data-testid="stDownloadButton"] > button,
    div[data-testid="stFormSubmitButton"] > button {
        font-size: 16px !important;
        font-weight: 800 !important;
        padding: 1.4em 2em !important;
        height: auto !important;
        border-radius: 14px !important;
        background-color: #90CAF9 !important;
        border-color: #90CAF9 !important;
        color: #0D3B66 !important;
    }
    div.stButton > button:hover,
    div.stDownloadButton > button:hover,
    div.stFormSubmitButton > button:hover,
    div[data-testid="stButton"] > button:hover,
    div[data-testid="stDownloadButton"] > button:hover,
    div[data-testid="stFormSubmitButton"] > button:hover {
        background-color: #64B5F6 !important;
        border-color: #64B5F6 !important;
        color: #0D3B66 !important;
    }
    div[data-testid="stHeading"] h3,
    div[data-testid="stMarkdownContainer"] h3,
    section.main h3 {
        font-size: 24px !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    '<h1 style="font-size:40px; margin:0;">📝 供應商帳單擷取核對系統</h1>',
    unsafe_allow_html=True,
)
st.markdown(
    """
    <div style="font-size:16px; color:#000000; font-weight:500; line-height:1.7; margin:12px 0 20px;">
        <div>① 上傳一份供應商帳單 PDF，系統自動辨識供應商並擷取欄位</div>
        <div>② 逐欄核對擷取結果：沒問題不用動作，有錯誤才按「❌ 標記為錯誤」並填入正確答案</div>
        <div>③ 核對完成後下載「回饋記錄 Excel」，可據此修正供應商帳單自動化辨識系統的擷取規則</div>
    </div>
    """,
    unsafe_allow_html=True,
)

with st.expander("ℹ️ 目前已支援的供應商清單"):
    for s in list_registered_suppliers():
        st.write(f"- **{s['label']}** (代碼: `{s['key']}`)")
    st.caption(
        "沒有列在上面的供應商，系統會顯示「未知供應商 (無法辨識)」，"
        "所有欄位都會是 N/A，這種情況也可以照樣核對、標記正確答案，"
        "回饋記錄可以做為之後新增這家供應商規則的素材。"
    )

if "uploader_version" not in st.session_state:
    st.session_state["uploader_version"] = 0

st.markdown('<span id="reset-btn-marker"></span>', unsafe_allow_html=True)
reset_clicked = st.button(
    "🔄 重新查核（換一份帳單）",
    help="清空目前的擷取結果、標記狀態，並清除已上傳的 PDF，方便重新上傳下一份帳單。",
)
if reset_clicked:
    st.session_state.pop("verify_result", None)
    st.session_state.pop("verify_signature", None)
    for k in list(st.session_state.keys()):
        if k.startswith("err_") or k.startswith("fix_"):
            del st.session_state[k]
    st.session_state["uploader_version"] += 1
    st.rerun()

# ---------------------------------------------------------------------------
# ① 上傳 PDF
# ---------------------------------------------------------------------------
st.subheader("① 上傳供應商帳單 PDF")
uploaded_pdf = st.file_uploader(
    "選擇一份 PDF 帳單",
    type=["pdf"],
    accept_multiple_files=False,
    key=f"pdf_upload_{st.session_state['uploader_version']}",
)

if not uploaded_pdf:
    st.info("請先上傳一份 PDF 帳單。")
    st.stop()

file_signature = hashlib.md5(uploaded_pdf.getvalue()).hexdigest()

run_clicked = st.button("▶️ 開始擷取", type="primary", use_container_width=True)

# 只有「換了新檔案」或「按下開始擷取」才重新跑一次擷取，避免使用者每次
# 勾選/輸入正確答案 (Streamlit 每次互動都會重新執行整支程式) 都重新跑一次
# 最多 20 次的擷取重試迴圈，浪費運算資源、也會讓畫面一直閃爍重置。
need_extract = (
    run_clicked
    or st.session_state.get("verify_signature") != file_signature
)

if need_extract:
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(uploaded_pdf.getvalue())
        tmp_path = tmp.name

    with st.spinner("解析中 (自動辨識供應商、嘗試不同擷取參數挑出資料最完整的一次)..."):
        quick_text = extract_text_from_pdf(tmp_path)
        supplier_key = detect_supplier(quick_text)
        result = find_best_extraction(tmp_path, supplier_key=supplier_key, reference_rows=None)

    st.session_state["verify_result"] = result
    st.session_state["verify_signature"] = file_signature
    # 換了新檔案，把之前殘留的標記狀態清掉，避免誤把上一份帳單的錯誤標記
    # 誤植到這一份帳單上。
    for k in list(st.session_state.keys()):
        if k.startswith("err_") or k.startswith("fix_"):
            del st.session_state[k]

result = st.session_state.get("verify_result")
if result is None:
    st.stop()

rows = result["rows"]
detected_key = result["supplier_key"]
quality = result["score"]

st.success(
    f"辨識供應商：**{supplier_label(detected_key)}** "
    f"（代碼：`{detected_key or '無法辨識'}`）　"
    f"資料完整度：**{quality * 100:.1f}%** "
    f"(欄位有抓到值、不是 N/A 的比例，僅供參考，實際對不對要靠下面人工核對)"
)

# ---------------------------------------------------------------------------
# ② 顯示擷取結果總覽
# ---------------------------------------------------------------------------
st.subheader("② 擷取結果總覽")
preview_df = pd.DataFrame(rows)[FIELD_CODES].rename(columns={c: DISPLAY_HEADERS[c] for c in FIELD_CODES})
st.dataframe(preview_df, use_container_width=True)

# ---------------------------------------------------------------------------
# ③ 逐欄人工核對
# ---------------------------------------------------------------------------
st.subheader("③ 逐欄人工核對（沒有標記 = 視為忽略，不會出現在回饋記錄）")

# 依 invoice_no 分組顯示：一般供應商只有一組 (一張帳單一個 invoice_no)，
# multi_page 供應商 (一份 PDF 可能塞好幾張獨立發票) 會自動依 invoice_no
# 分成好幾組，方便使用者知道現在核對的是哪一張發票。
invoice_groups: "list[tuple[str, list[int]]]" = []
seen_invoice = {}
for idx, row in enumerate(rows):
    inv_no = row.get("invoice_no") or NA
    seen_invoice.setdefault(inv_no, []).append(idx)
for inv_no, idxs in seen_invoice.items():
    invoice_groups.append((inv_no, idxs))

marked_count = 0
total_fields = 0

for inv_no, idxs in invoice_groups:
    group_label = f"發票號碼：{inv_no}" if len(invoice_groups) > 1 else "帳單內容"
    st.markdown(f"#### 📄 {group_label}")
    for row_idx in idxs:
        row = rows[row_idx]
        desc = row.get("description", NA)
        amount = row.get("amount", NA)
        with st.expander(f"品項 {row_idx + 1}：{desc}　|　金額：{amount}", expanded=(len(rows) <= 3)):
            for code in FIELD_CODES:
                total_fields += 1
                value = row.get(code, NA)
                label = DISPLAY_HEADERS[code].split("\n")[0]
                col_label, col_value, col_flag, col_fix = st.columns([2, 3, 2, 3])
                col_label.markdown(f"**{label}**")
                col_value.write(value)
                err_key = f"err_{file_signature}_{row_idx}_{code}"
                fix_key = f"fix_{file_signature}_{row_idx}_{code}"
                marked = col_flag.checkbox("❌ 標記為錯誤", key=err_key)
                if marked:
                    marked_count += 1
                    col_fix.text_input(
                        "✏️ 正確答案（留空代表正確答案應為 N/A）",
                        key=fix_key,
                        label_visibility="visible",
                    )
                else:
                    col_fix.write("")

st.info(f"目前已標記 **{marked_count}** / {total_fields} 個欄位為錯誤，其餘視為忽略。")

# ---------------------------------------------------------------------------
# ④ 產生回饋記錄 Excel
# ---------------------------------------------------------------------------
st.subheader("④ 下載回饋記錄")


def _build_feedback_excel(file_name: str, supplier_lbl: str, supplier_key_: str,
                           rows_: list, marks: dict) -> bytes:
    """把「被標記為錯誤的欄位 + 使用者填的正確答案」整理成一份 Excel，
    欄位：檔案名稱／供應商／發票號碼／品項序號／欄位代碼／欄位名稱／
    系統擷取值／人工輸入的正確答案，方便日後依此修正 suppliers.py。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "回饋記錄"

    headers = ["檔案名稱", "供應商", "供應商代碼", "發票號碼", "品項序號",
               "欄位代碼", "欄位名稱", "系統擷取值", "人工輸入正確答案"]
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = 20

    r = 2
    for row_idx, row_ in enumerate(rows_):
        for code in FIELD_CODES:
            err_key = f"err_{file_signature}_{row_idx}_{code}"
            fix_key = f"fix_{file_signature}_{row_idx}_{code}"
            if not marks.get(err_key):
                continue
            correct_value = marks.get(fix_key, "") or "N/A"
            ws.cell(row=r, column=1, value=file_name)
            ws.cell(row=r, column=2, value=supplier_lbl)
            ws.cell(row=r, column=3, value=supplier_key_ or "")
            ws.cell(row=r, column=4, value=row_.get("invoice_no", NA))
            ws.cell(row=r, column=5, value=row_idx + 1)
            ws.cell(row=r, column=6, value=code)
            ws.cell(row=r, column=7, value=DISPLAY_HEADERS[code].split("\n")[0])
            ws.cell(row=r, column=8, value=row_.get(code, NA))
            ws.cell(row=r, column=9, value=correct_value)
            r += 1

    if r == 2:
        ws.cell(row=2, column=1, value="（目前沒有任何欄位被標記為錯誤）")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if marked_count == 0:
    st.caption("目前沒有任何欄位被標記為錯誤，標記完成後即可在這裡下載回饋記錄 Excel。")
else:
    feedback_bytes = _build_feedback_excel(
        uploaded_pdf.name, supplier_label(detected_key), detected_key, rows, st.session_state
    )
    st.download_button(
        "⬇️ 下載回饋記錄 (Excel)",
        data=feedback_bytes,
        file_name=f"回饋記錄_{uploaded_pdf.name.rsplit('.', 1)[0]}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    st.caption(
        "把這份 Excel 交給負責維護「供應商帳單自動化辨識系統」的開發端 "
        "(或直接請 Claude 依照 add-invoice-supplier 的流程) 修正對應供應商的擷取規則，"
        "改完後回到這個網頁重新上傳同一份 PDF，確認錯誤欄位已修正。"
    )
