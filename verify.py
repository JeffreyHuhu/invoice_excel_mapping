# -*- coding: utf-8 -*-
"""
verify.py
=========
Step 2：查核比對小程式

目的：
  Step1 (app.py) 把 PDF 轉成 Excel 之後，這支程式會「獨立地」重新解析
  同一批 PDF 原始檔，把重新解析出來的結果拿去和 Excel 裡的資料逐欄比對，
  找出兩邊不一致的地方，並計算每個欄位、以及整體的「正確率」。

  這樣就不用再靠人工逐筆核對 PDF 與 Excel，改成程式自動比對、
  只需要人工複查被標記為「不一致」的少數欄位即可，大幅減少人工工時。

使用方式 (命令列)：
    python verify.py --pdf_dir ./pdfs --excel_path ./DWIHARTA_轉檔結果.xlsx \
                      --report_path ./查核報告.xlsx

參數說明：
    --pdf_dir      放置原始 PDF 帳單的資料夾 (會遞迴搜尋所有 .pdf)
    --excel_path   Step1 產出 (或人工複核後) 的 Excel 檔案
    --report_path  查核結果報告要輸出的路徑 (Excel)

比對邏輯：
  1. 用 Invoice No 把 Excel 中的資料列 group 起來 (同一張帳單的多筆費用列)
  2. 針對每個 Invoice No，在 pdf_dir 底下重新解析出「基準答案」(ground truth)
  3. 抬頭欄位 (Invoice Date、B/L No... 等) 只需比對一次
  4. 費用明細 (Description / Amount) 用「集合比對」：
     - 用 (description 正規化後的字串, 金額) 當作一個費用項目的 key
     - 比較 Excel 中出現的費用項目集合 與 PDF 重新解析出的集合是否一致
  5. 輸出:
       - 逐欄比對結果 (相符 / 不相符 / Excel缺漏 / PDF缺漏)
       - 每個欄位的正確率統計
       - 整體正確率 (相符欄位數 / 應比對欄位總數)
"""

import argparse
import glob
import os
import re
import sys
import datetime
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font

from extract_utils import (
    parse_invoice_pdf,
    TEMPLATE_COLUMNS,
    normalize_value as normalize,
    compare_invoice,
    compute_accuracy,
)

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")


# ---------------------------------------------------------------------------
# 工具函式 (normalize / compare_invoice 現在改從 extract_utils 共用，
# 避免這支獨立 CLI 工具跟 app.py 的比對邏輯日後跑掉、兜不起來)
# ---------------------------------------------------------------------------

def load_pdf_ground_truth(pdf_dir: str) -> dict:
    """重新解析 pdf_dir 底下所有 PDF，回傳 {invoice_no: ParsedInvoice}。"""
    ground_truth = {}
    pdf_paths = glob.glob(os.path.join(pdf_dir, "**", "*.pdf"), recursive=True)
    if not pdf_paths:
        print(f"⚠️  在 {pdf_dir} 找不到任何 PDF 檔案", file=sys.stderr)
    for path in pdf_paths:
        try:
            parsed = parse_invoice_pdf(path)
            invoice_no = normalize(parsed.header.invoice_no)
            if not invoice_no:
                print(f"⚠️  {path} 抓不到 Invoice No，略過", file=sys.stderr)
                continue
            ground_truth[invoice_no] = parsed
        except Exception as e:  # noqa: BLE001
            print(f"❌  解析 {path} 失敗：{e}", file=sys.stderr)
    return ground_truth


def load_excel_rows(excel_path: str) -> pd.DataFrame:
    """讀取 Step1 產出的 Excel (資料從第3列開始，欄位對應 TEMPLATE_COLUMNS)。"""
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    records = []
    for row in ws.iter_rows(min_row=3, values_only=False):
        values = {}
        empty = True
        for idx, col_name in enumerate(TEMPLATE_COLUMNS, start=2):  # B欄起
            cell = row[idx - 1]
            values[col_name] = cell.value
            if cell.value not in (None, ""):
                empty = False
        if not empty:
            records.append(values)
    return pd.DataFrame(records)


# ---------------------------------------------------------------------------
# 比對邏輯
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# 比對邏輯 (compare_invoice 已搬到 extract_utils.py，這裡直接沿用)
# ---------------------------------------------------------------------------

def run_verification(pdf_dir: str, excel_path: str) -> pd.DataFrame:
    ground_truth = load_pdf_ground_truth(pdf_dir)
    excel_df = load_excel_rows(excel_path)

    if excel_df.empty:
        raise ValueError("Excel 檔案中沒有讀到任何資料列，請確認格式是否正確。")

    all_results = []
    for invoice_no, group in excel_df.groupby(excel_df["Invoice No"].map(normalize)):
        parsed_pdf = ground_truth.get(invoice_no)
        if parsed_pdf is None:
            all_results.append({
                "Invoice No": invoice_no,
                "欄位": "(整張帳單)",
                "正確答案值": invoice_no,
                "擷取結果值": "(找不到對應的PDF檔案)",
                "結果": "不相符",
            })
            continue
        for rec in compare_invoice(group, parsed_pdf):
            rec_with_id = {"Invoice No": invoice_no, **rec}
            all_results.append(rec_with_id)

    return pd.DataFrame(all_results)


def summarize(result_df: pd.DataFrame) -> pd.DataFrame:
    """依「欄位」彙總正確率。"""
    summary = (
        result_df.groupby("欄位")["結果"]
        .apply(lambda s: pd.Series({
            "比對筆數": len(s),
            "相符筆數": (s == "相符").sum(),
            "正確率": f"{(s == '相符').mean() * 100:.1f}%",
        }))
        .unstack()
    )
    summary = summary[["比對筆數", "相符筆數", "正確率"]]
    total = len(result_df)
    matched = (result_df["結果"] == "相符").sum()
    overall = pd.DataFrame(
        [{"比對筆數": total, "相符筆數": matched,
          "正確率": f"{matched / total * 100:.1f}%" if total else "N/A"}],
        index=["整體 (Overall)"],
    )
    return pd.concat([summary, overall])


def write_report(result_df: pd.DataFrame, summary_df: pd.DataFrame, report_path: str):
    wb = Workbook()

    # --- Sheet1: 正確率統計摘要 ---
    ws1 = wb.active
    ws1.title = "正確率摘要"
    ws1.append(["欄位", "比對筆數", "相符筆數", "正確率"])
    for cell in ws1[1]:
        cell.font = Font(bold=True)
    for idx, row in summary_df.iterrows():
        ws1.append([idx, int(row["比對筆數"]), int(row["相符筆數"]), row["正確率"]])
    for col_letter, width in zip("ABCD", (22, 12, 12, 12)):
        ws1.column_dimensions[col_letter].width = width

    # --- Sheet2: 逐筆比對明細 (不相符標紅、相符標綠) ---
    ws2 = wb.create_sheet("逐筆比對明細")
    cols = ["Invoice No", "欄位", "正確答案值", "擷取結果值", "結果"]
    ws2.append(cols)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for _, r in result_df.iterrows():
        ws2.append([r[c] for c in cols])
        fill = GREEN if r["結果"] == "相符" else RED
        ws2.cell(row=ws2.max_row, column=5).fill = fill
    widths = (18, 16, 30, 30, 10)
    for col_letter, width in zip("ABCDE", widths):
        ws2.column_dimensions[col_letter].width = width

    wb.save(report_path)


# ---------------------------------------------------------------------------
# CLI 入口
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="查核比對 PDF 帳單與 Excel 轉檔結果的正確率")
    parser.add_argument("--pdf_dir", required=True, help="原始 PDF 帳單資料夾")
    parser.add_argument("--excel_path", required=True, help="Step1 產出的 Excel 檔案路徑")
    parser.add_argument("--report_path", default="查核報告.xlsx", help="查核報告輸出路徑")
    args = parser.parse_args()

    print("🔍 重新解析 PDF 並與 Excel 比對中...")
    result_df = run_verification(args.pdf_dir, args.excel_path)
    summary_df = summarize(result_df)

    print("\n===== 正確率摘要 =====")
    print(summary_df.to_string())

    write_report(result_df, summary_df, args.report_path)
    print(f"\n✅ 詳細查核報告已輸出至：{args.report_path}")


if __name__ == "__main__":
    main()
