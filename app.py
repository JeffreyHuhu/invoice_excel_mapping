# -*- coding: utf-8 -*-
"""
app.py
======
供應商 PDF 帳單自動化系統 (轉檔 + 自動查核正確率，單一流程)

跟前一版的差異：
  - 不再需要「上傳PDF」+「另外再上傳一份填好的Excel」兩次操作才能查核。
  - 上傳 PDF 後，程式對每一份 PDF 用多組擷取參數 (EXTRACTION_CONFIGS) 各
    解析一次，並用迴圈挑出分數最高的一組結果 -> 這就是 Step1 的轉檔結果。
  - 轉檔完成後，正確率會直接顯示在同一個頁面上，不用再手動觸發查核。
  - 若你手邊剛好已經有人工核對過的正確答案 Excel，也可以「選填」上傳，
    這樣顯示的就是跟正確答案逐欄比對過的『真正正確率』；沒有上傳的話，
    顯示的是『資料品質分數』(完整度+是否有欄位互相污染)，兩者在頁面上
    會清楚分開標示，不會混為一談。
"""

import io
import datetime
import tempfile
import os

import streamlit as st
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from extract_utils import (
    TEMPLATE_COLUMNS,
    invoice_to_rows,
    parse_invoice_pdf,
    parse_invoice_pdf_best,
    compare_invoice,
    compute_accuracy,
    normalize_value,
)

# ---------------------------------------------------------------------------
# 樣式與對照設定
# ---------------------------------------------------------------------------

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")

DISPLAY_HEADERS = {
    "Invoice Date": "INVOICE DATE\n發票日期",
    "Invoice No": "INVOICE NO\n發票號碼",
    "Payment can be Transfer to": "SUPPLIER\n供應商名稱",
    "TO": "CONSIGNEE\n集團公司名稱",
    "Description\nVAT": "Description\n費用名稱",
    "Amount (IDR": "AMOUNT\n金額",
    "Order No.": "Order No.",
    "Arrive Date": "Arrive Date\n到達日",
    "on board date": "on board date\n開船日",
    "B/L NO.": "B/L NO.\n提單號碼",
    "MBL NO.": "MBL NO.\n主提單號碼",
    "Port of Loading": "Port of Loading\n啟運港",
    "Port of discharge": "Port of discharge\n目的港",
    "Volume": "Volume\n材積",
    "Vessel": "Vessel\n船名",
    "Container no.": "Container no.\n貨櫃號碼",
}


# ---------------------------------------------------------------------------
# 共用輔助函式
# ---------------------------------------------------------------------------

def build_excel(df: pd.DataFrame) -> bytes:
    """把整理好的 DataFrame 輸出成符合範本格式的 Excel (bytes)。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "DWIHARTA"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    wrap_center = Alignment(wrap_text=True, vertical="center", horizontal="center")

    ws.cell(row=1, column=1, value="抓取欄位")
    ws.cell(row=2, column=1, value="顯示欄位")

    for idx, col_name in enumerate(TEMPLATE_COLUMNS, start=2):
        ws.cell(row=1, column=idx, value=col_name)
        cell2 = ws.cell(row=2, column=idx, value=DISPLAY_HEADERS.get(col_name, col_name))
        cell2.font = bold
        cell2.fill = header_fill
        cell2.alignment = wrap_center
        ws.column_dimensions[get_column_letter(idx)].width = 20

    for r, (_, row) in enumerate(df.iterrows(), start=3):
        for idx, col_name in enumerate(TEMPLATE_COLUMNS, start=2):
            value = row.get(col_name, "")
            if isinstance(value, (datetime.date, datetime.datetime)):
                cell = ws.cell(row=r, column=idx, value=value)
                cell.number_format = "yyyy-mm-dd"
            else:
                ws.cell(row=r, column=idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def load_excel_rows_from_bytes(excel_bytes: bytes) -> pd.DataFrame:
    """從上傳的「正確答案」Excel bytes 讀取資料列 (格式需與範本一致)。"""
    wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
    ws = wb.active
    records = []
    for row in ws.iter_rows(min_row=3, values_only=False):
        values = {}
        empty = True
        for idx, col_name in enumerate(TEMPLATE_COLUMNS, start=2):
            cell = row[idx - 1]
            values[col_name] = cell.value
            if cell.value not in (None, ""):
                empty = False
        if not empty:
            records.append(values)
    return pd.DataFrame(records)


def generate_report_bytes(result_df: pd.DataFrame, summary_df: pd.DataFrame) -> bytes:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "正確率摘要"
    ws1.append(["欄位", "比對筆數", "相符筆數", "正確率"])
    for cell in ws1[1]:
        cell.font = Font(bold=True)
    for idx, row in summary_df.iterrows():
        ws1.append([idx, int(row["比對筆數"]), int(row["相符筆數"]), row["正確率"]])
    for col_letter, width in zip("ABCD", (22, 12, 12, 12)):
        ws1.column_dimensions[col_letter].width = width

    ws2 = wb.create_sheet("逐筆比對明細")
    cols = ["Invoice No", "欄位", "正確答案值", "擷取結果值", "結果"]
    ws2.append(cols)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for _, r in result_df.iterrows():
        ws2.append([r[c] for c in cols])
        fill = GREEN if r["結果"] == "相符" else RED
        ws2.cell(row=ws2.max_row, column=5).fill = fill
    widths = (18, 16, 30, 30, 10)
    for col_letter, width in zip("ABCDE", widths):
        ws2.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def summarize(result_df: pd.DataFrame) -> pd.DataFrame:
    summary = (
        result_df.groupby("欄位")["結果"]
        .apply(lambda s: pd.Series({
            "比對筆數": len(s),
            "相符筆數": (s == "相符").sum(),
            "正確率": f"{(s == '相符').mean() * 100:.1f}%",
        }))
        .unstack()
    )
    summary = summary[["比對筆數", "相符筆數", "正確率"]]
    total = len(result_df)
    matched = (result_df["結果"] == "相符").sum()
    overall = pd.DataFrame(
        [{"比對筆數": total, "相符筆數": matched,
          "正確率": f"{matched / total * 100:.1f}%" if total else "N/A"}],
        index=["整體 (Overall)"],
    )
    return pd.concat([summary, overall])


# ---------------------------------------------------------------------------
# Streamlit 網頁主介面
# ---------------------------------------------------------------------------

st.set_page_config(page_title="供應商帳單自動化與查核系統", layout="wide")
st.title("📄➡️📊 供應商帳單自動化系統 (轉檔＋自動查核正確率)")
st.caption(
    "上傳 PDF 後，系統會用多組擷取設定各解析一次 (迴圈)，自動挑出分數最高的結果，"
    "並直接在下方顯示轉檔正確率，不需要再手動上傳一次 Excel 做比對。"
)

col_pdf, col_ref = st.columns(2)
with col_pdf:
    uploaded_pdfs = st.file_uploader(
        "① 上傳 PDF 帳單 (可一次選取多個檔案)",
        type=["pdf"],
        accept_multiple_files=True,
        key="pdf_upload",
    )
with col_ref:
    reference_excel = st.file_uploader(
        "② (選填) 若手邊已有人工核對過的正確答案 Excel，可上傳取得『真正正確率』",
        type=["xlsx"],
        accept_multiple_files=False,
        key="reference_upload",
    )
    st.caption("沒有上傳的話，會改用『資料品質分數』(欄位完整度＋是否有互相污染) 作為代理指標。")

if uploaded_pdfs:
    reference_df = None
    if reference_excel is not None:
        try:
            reference_df = load_excel_rows_from_bytes(reference_excel.getvalue())
            if reference_df.empty:
                st.warning("⚠️ 上傳的正確答案 Excel 讀不到任何資料列，將改用資料品質分數。")
                reference_df = None
        except Exception as e:  # noqa: BLE001
            st.warning(f"⚠️ 正確答案 Excel 讀取失敗 ({e})，將改用資料品質分數。")

    has_reference = reference_df is not None
    score_label = "正確率" if has_reference else "品質分數"

    all_rows = []
    parse_errors = []
    per_invoice_scores = []          # 每份 PDF 的最佳分數 (給頁面上的總覽表)
    per_invoice_results = []          # 有正確答案時，逐欄比對明細 (給下載報告用)
    config_attempts_log = []          # 每份 PDF 迴圈嘗試各組設定的分數紀錄 (可攤開查看)

    with st.spinner("解析中 (每份 PDF 會嘗試多組設定，挑出分數最高的結果)..."):
        for f in uploaded_pdfs:
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                tmp.write(f.getvalue())
                tmp_path = tmp.name
            try:
                # 若有提供正確答案 Excel，先用預設設定快速解析拿到 Invoice No，
                # 藉此在正確答案 Excel 裡找出「這份 PDF 對應的那幾列」，
                # 之後迴圈評分時才能比對到正確的那一張帳單。
                ref_group = None
                if has_reference:
                    quick = parse_invoice_pdf(tmp_path)
                    inv_no = normalize_value(quick.header.invoice_no)
                    if inv_no:
                        mask = reference_df["Invoice No"].map(normalize_value) == inv_no
                        candidate = reference_df[mask]
                        if not candidate.empty:
                            ref_group = candidate

                best_parsed, best_score, best_cfg, attempts = parse_invoice_pdf_best(
                    tmp_path, reference_group=ref_group
                )

                if not best_parsed.line_items:
                    parse_errors.append(f"⚠️ {f.name}：找不到任何費用明細，請確認 PDF 是否為可編輯文字檔。")

                rows = invoice_to_rows(best_parsed)
                for row in rows:
                    row["__來源檔案"] = f.name
                all_rows.extend(rows)

                per_invoice_scores.append({
                    "來源檔案": f.name,
                    "Invoice No": best_parsed.header.invoice_no,
                    score_label: f"{best_score * 100:.1f}%",
                    "採用設定": best_cfg or "預設",
                    "比對基準": "正確答案 Excel" if ref_group is not None else "資料品質自我檢查",
                })
                config_attempts_log.append({"檔案": f.name, "嘗試紀錄": attempts})

                if ref_group is not None:
                    results = compare_invoice(ref_group, best_parsed)
                    for r in results:
                        r["Invoice No"] = best_parsed.header.invoice_no
                    per_invoice_results.extend(results)

            except Exception as e:  # noqa: BLE001
                parse_errors.append(f"❌ {f.name}：解析失敗 ({e})")
            finally:
                os.remove(tmp_path)

    if parse_errors:
        for msg in parse_errors:
            st.warning(msg)

    if all_rows:
        df = pd.DataFrame(all_rows)
        display_cols = ["__來源檔案"] + TEMPLATE_COLUMNS
        df = df[display_cols].rename(columns={"__來源檔案": "來源檔案"})

        st.subheader("✏️ 轉檔結果預覽（可直接在表格中修正錯誤欄位）")
        edited_df = st.data_editor(df, num_rows="dynamic", use_container_width=True)

        export_df = edited_df.rename(columns={"來源檔案": "__來源檔案"})
        export_df = export_df[TEMPLATE_COLUMNS]
        excel_bytes = build_excel(export_df)

        st.download_button(
            label="⬇️ 下載 Excel (符合 DWIHARTA 範本格式)",
            data=excel_bytes,
            file_name=f"DWIHARTA_轉檔結果_{datetime.date.today().isoformat()}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        # --- 自動顯示轉檔正確率 ---
        st.divider()
        st.subheader(f"📊 轉檔{score_label}")
        if not has_reference:
            st.info(
                "目前沒有上傳正確答案 Excel，以下是『資料品質分數』：檢查每個欄位是否"
                "有抓到值、且沒有把下一個欄位的文字誤黏進來。這不是跟人工核對過的正確率，"
                "但能反映轉檔是否乾淨；建議日後累積幾份人工核對過的 Excel 後上傳，取得真正正確率。"
            )

        score_df = pd.DataFrame(per_invoice_scores)
        st.dataframe(score_df, use_container_width=True)

        # 平均分數 (把百分比字串轉回數字算平均)
        avg_score = score_df[score_label].str.rstrip("%").astype(float).mean()
        st.metric(f"整體轉檔{score_label} (所有 PDF 平均)", f"{avg_score:.1f}%")

        with st.expander("🔍 查看每份 PDF 在迴圈中嘗試過的各組擷取設定分數"):
            for log in config_attempts_log:
                st.write(f"**{log['檔案']}**")
                st.dataframe(pd.DataFrame(log["嘗試紀錄"]), use_container_width=True)

        if has_reference and per_invoice_results:
            result_df = pd.DataFrame(per_invoice_results)
            summary_df = summarize(result_df)
            st.subheader("📋 與正確答案逐欄比對明細（依欄位分類的正確率）")
            st.dataframe(summary_df, use_container_width=True)

            report_bytes = generate_report_bytes(result_df, summary_df)
            st.download_button(
                label="⬇️ 下載詳細紅綠燈查核報告 (Excel)",
                data=report_bytes,
                file_name=f"查核報告_{datetime.date.today().isoformat()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
else:
    st.info("請先上傳至少一份 PDF 帳單。")
